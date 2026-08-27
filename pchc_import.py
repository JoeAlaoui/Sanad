# -*- coding: utf-8 -*-
"""
pchc_import.py — Parsing des exports "Dossiers PCHC" (ex-reporting Tarkhiss dossiers).

Fonctions pures : aucune I/O disque, aucun appel Flask. Reçoivent un classeur/fichier et
retournent des enregistrements normalisés, à charge de l'appelant (app.py) de les stocker.
"""
import re
import csv as _csv
import io as _io
from datetime import datetime

# Catégories gérées : clé interne -> (libellé, en-tête feuille attendue, colonne "entité", colonne "libellé produit")
CATEGORIES = {
    "identification": {
        "label": "Identification Opérateur",
        "sheet_match": ["identification", "operateur", "opérateur"],
        "entity_col": ["Raison sociale"],
        "detail_col": ["Type opérateur"],
        "ref_col": ["ID"],
        "date_col": ["Date modification", "Date dépôt"],
    },
    "declaration": {
        "label": "Déclaration d'Activité",
        "sheet_match": ["declaration", "déclaration"],
        "entity_col": ["Etablissement"],
        "detail_col": ["Type demande", "Type produit"],
        "ref_col": ["Réf. dossier", "Réf.Dossier"],
        "date_col": ["Date dépôt"],
    },
    "ce": {
        "label": "Certificat d'Enregistrement",
        "sheet_match": ["certificatenregistrement", "certificat enreg"],
        "entity_col": ["Etablissement"],
        "detail_col": ["Nom produit", "Type demande"],
        "ref_col": ["Réf.Dossier", "Réf. dossier"],
        "date_col": ["Date dépôt"],
    },
    "clv": {
        "label": "Certificat de Libre Vente",
        "sheet_match": ["certificatlibrevente", "libre vente"],
        "entity_col": ["Etablissement"],
        "detail_col": ["Nom produit"],
        "ref_col": ["Réf. dossier", "Réf.Dossier"],
        "date_col": ["Date dépôt"],
    },
    "aimp": {
        "label": "Autorisation d'Importation Matières Premières",
        "sheet_match": ["autorisationimportation", "matiere", "matière", "aimp"],
        "entity_col": ["Importateur"],
        "detail_col": [],
        "ref_col": ["Réf.Dossier", "Réf. dossier"],
        "date_col": ["Date dépôt"],
    },
}

# Table de correspondance statut -> couleur par défaut (clé normalisée, sans accents/casse)
DEFAULT_STATUS_COLORS = {
    "delivre": "green", "valide": "green", "identifie": "green",
    "en cours": "blue", "en evaluation": "blue", "en verification": "blue",
    "validation acte": "blue", "edition acte": "blue", "evaluation complements": "blue",
    "attente validation evaluation": "blue", "seance de travail": "blue", "signature acte": "blue",
    "non conforme": "yellow",
    "attente paiement": "orange", "attente signature": "orange",
    "irrecevable": "red", "rejete": "red",
}

COLOR_HEX = {
    "green": "#25935F", "blue": "#1794CF", "yellow": "#F2C94C",
    "orange": "#F2994A", "red": "#DC2828", "grey": "#94A3AD",
}


def normalize_key(s):
    if not s:
        return ""
    s = str(s).strip().lower()
    repl = {"é": "e", "è": "e", "ê": "e", "à": "a", "â": "a", "ô": "o", "î": "i",
            "ç": "c", "ù": "u", "û": "u", "ë": "e", "ï": "i"}
    for a, b in repl.items():
        s = s.replace(a, b)
    return s


def status_color(statut, custom_map=None):
    key = normalize_key(statut)
    m = dict(DEFAULT_STATUS_COLORS)
    if custom_map:
        m.update({normalize_key(k): v for k, v in custom_map.items()})
    return m.get(key, "grey")


def parse_date_any(value):
    """Parse une date au format ISO datetime, 'YYYY-MM-DD HH:MM:SS' ou 'DD/MM/YYYY'."""
    if value is None or value == "":
        return None
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d")
    s = str(value).strip()
    if not s:
        return None
    m = re.match(r"^(\d{4})-(\d{2})-(\d{2})", s)
    if m:
        return f"{m.group(1)}-{m.group(2)}-{m.group(3)}"
    m = re.match(r"^(\d{1,2})/(\d{1,2})/(\d{4})", s)
    if m:
        d, mo, y = m.group(1), m.group(2), m.group(3)
        return f"{y}-{mo.zfill(2)}-{d.zfill(2)}"
    return None


def find_col(headers, candidates):
    norm_headers = [normalize_key(h) for h in headers]
    for cand in candidates:
        nc = normalize_key(cand)
        for i, h in enumerate(norm_headers):
            if nc == h or nc in h:
                return i
    return None


def match_category(sheet_name):
    norm = normalize_key(sheet_name)
    for key, cfg in CATEGORIES.items():
        for m in cfg["sheet_match"]:
            if normalize_key(m) in norm:
                return key
    return None


def parse_category_rows(rows, headers, cat_key):
    cfg = CATEGORIES[cat_key]
    idx_ref = find_col(headers, cfg["ref_col"])
    idx_entity = find_col(headers, cfg["entity_col"])
    idx_statut = find_col(headers, ["Statut"])
    idx_date = None
    for dc in cfg["date_col"]:
        idx_date = find_col(headers, [dc])
        if idx_date is not None:
            break
    idx_details = [find_col(headers, [d]) for d in cfg["detail_col"]]
    idx_details = [i for i in idx_details if i is not None]

    records = []
    for row in rows:
        if row is None or (idx_ref is not None and not row[idx_ref]):
            continue
        ref = str(row[idx_ref]).strip() if idx_ref is not None and row[idx_ref] else ""
        if not ref:
            continue
        entity = str(row[idx_entity]).strip() if idx_entity is not None and row[idx_entity] else ""
        statut = str(row[idx_statut]).strip() if idx_statut is not None and row[idx_statut] else ""
        date_depot = parse_date_any(row[idx_date]) if idx_date is not None else None
        details = " — ".join(str(row[i]).strip() for i in idx_details if row[i]) if idx_details else ""
        records.append({
            "ref": ref, "entity": entity, "statut": statut,
            "date_depot": date_depot, "details": details,
        })
    return records


def import_pchc_xlsx(file_storage):
    """Parse un classeur consolidé (.xlsx). Retourne {cat_key: [records]}."""
    from openpyxl import load_workbook
    wb = load_workbook(file_storage, data_only=True)
    out = {}
    for sheet_name in wb.sheetnames:
        cat_key = match_category(sheet_name)
        if not cat_key:
            continue
        ws = wb[sheet_name]
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            continue
        headers = [str(h) if h is not None else "" for h in rows[0]]
        out[cat_key] = parse_category_rows(rows[1:], headers, cat_key)
    return out


def import_pchc_csv(file_storage, cat_key):
    """Parse un export CSV séparé pour une catégorie donnée. Retourne [records]."""
    text = file_storage.read().decode("utf-8-sig", errors="replace")
    delim = ";" if ";" in text.split("\n")[0] else ","
    reader = _csv.reader(_io.StringIO(text), delimiter=delim)
    rows = list(reader)
    if not rows:
        return []
    headers = rows[0]
    return parse_category_rows(rows[1:], headers, cat_key)
